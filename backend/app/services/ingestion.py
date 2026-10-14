import re
from pathlib import Path
from sqlalchemy import select

from ..config import settings
from ..models import Document, Knowledge, KnowledgeBase, ModelProfile
from ..schemas import KnowledgeInput, ParseOptions
from .knowledge import touch
from .splitting import split_blocks
from .recognition import RECOGNIZABLE, run_recognition


def parse_blocks(path: Path, filename: str):
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if settings().semantica_enabled:
            from semantica.parse import DocumentParser
            parsed = DocumentParser().parse(str(path), extract_tables=False, extract_images=False)
            return [(f"PDF 第 {page['page_number']} 页", page.get("text", "")) for page in parsed["pages"]]
        from pypdf import PdfReader
        return [(f"PDF 第 {i+1} 页", p.extract_text() or "") for i, p in enumerate(PdfReader(path).pages)]
    if suffix == ".pptx":
        from pptx import Presentation
        return [(f"幻灯片 {i+1}", "\n".join(s.text for s in slide.shapes if s.has_text_frame))
                for i, slide in enumerate(Presentation(path).slides)]
    if suffix == ".docx":
        from docx import Document as WordDocument
        doc = WordDocument(path)
        return [(f"段落 {i+1}", p.text) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
    if suffix in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8-sig")
        sections = re.split(r"(?m)(?=^#{1,3} )", text)
        return [(f"章节 {i+1}：{s.splitlines()[0][:80]}", s) for i, s in enumerate(sections) if s.strip()]
    raise ValueError("仅支持 PDF、PPTX、DOCX、Markdown 和 TXT")


def resolve_split_profile(db, options, kb):
    if options.splitting.method != "semantic":
        return None
    profile_id = options.splitting.model_profile_id or kb.extraction_profile_id
    profile = db.get(ModelProfile, profile_id) if profile_id else None
    if not profile or profile.provider == "demo" or profile.purpose != "generation":
        raise ValueError("模型语义切分必须选择真实的生成模型配置")
    options.splitting.model_profile_id = profile.id
    return profile


def ingest(db, document_id, actor_id, progress=None):
    doc = db.get(Document, document_id)
    if not doc or doc.withdrawn:
        raise ValueError("文档不存在或已撤回")
    kb = db.get(KnowledgeBase, doc.kb_id)
    options = ParseOptions.model_validate(doc.parse_options or {})
    profile = resolve_split_profile(db, options, kb)
    if doc.state in {"ready", "review"} and (doc.recognition or {}).get("splitting"):
        existing = [item for item in db.scalars(select(Knowledge).where(Knowledge.kb_id == kb.id, ~Knowledge.withdrawn))
                    if any(s.get("document_id") == doc.id for s in item.sources)]
        if existing:
            doc.state, doc.parse_options = "ready", options.model_dump()
            db.commit()
            return {"knowledge_count": len(existing), "document_id": doc.id, "splitting": doc.recognition["splitting"]}
    suffix = Path(doc.filename).suffix.lower()
    # PDF and images always pass through OCR, including legacy mode=text requests.
    if suffix in RECOGNIZABLE:
        if not settings().ocr_enabled:
            raise ValueError("PDF 与图片必须先进行 OCR，请启用 PaddleOCR 后重试")
        report = run_recognition(doc, progress=(lambda value: progress(10+int(value*.35))) if progress else None)
        blocks, parser = report["blocks"], report["parser"]
        doc.recognition = {k: report[k] for k in ["parser", "counts", "warnings"]}
    else:
        if options.mode == "scan":
            raise ValueError("Office 图文识别请先导出 PDF；文字文档请使用自动模式")
        blocks = [{"locator": where, "text": text, "label": "text", "warnings": [],
                   **({"page": i+1} if suffix == ".pptx" else {})}
                  for i, (where, text) in enumerate(parse_blocks(Path(doc.path), doc.filename))]
        parser = "native-text"
    chunks, splitting = split_blocks(blocks, options.splitting.model_dump(), profile,
        cache_dir=settings().storage_dir/"splitting"/doc.id/"cache",
        progress=(lambda ratio: progress(45+int(ratio*50))) if progress else None)
    created = []
    for index, chunk in enumerate(chunks):
        sources = []
        for source in chunk["sources"]:
            block = blocks[source["block_index"]]
            sources.append({"document_id": doc.id, "locator": block["locator"], "excerpt": source["excerpt"],
                "note": "保留原文识别结果与来源位置，可对照原图编辑", "parser": parser,
                **{k: block.get(k) for k in ["page", "bbox", "asset_id", "confidence"]},
                "warnings": block.get("warnings", [])})
        attributes = {"chapter": sources[0]["locator"], "splitting": {
            "method": options.splitting.method, "model_profile_id": profile.id if profile else None,
            "model": profile.model if profile else None, "chunk_index": index,
            "range": [chunk["start"], chunk["end"]], "source_block_indices": [s["block_index"] for s in chunk["sources"]]},
            "extraction": "模型语义切分及知识类型标注；未自动构建知识关系" if profile else "程序规则切分；未进行模型语义提取"}
        if chunk["kind"] == "formula" and len(chunk["sources"]) == 1:
            attributes["latex"] = blocks[chunk["sources"][0]["block_index"]].get("latex", "")
        parsed = KnowledgeInput(kind=chunk["kind"], title=chunk["title"], content=chunk["content"],
            sources=sources, attributes=attributes)
        item = Knowledge(kb_id=kb.id, **parsed.model_dump(exclude={"change_note"}))
        db.add(item)
        created.append(item)
    db.refresh(doc, attribute_names=["withdrawn"])
    if doc.withdrawn:
        raise ValueError("处理期间文档已撤回，未导入新知识")
    doc.state = "ready"
    doc.parse_options = options.model_dump()
    doc.recognition = {**(doc.recognition or {}), "splitting": splitting}
    touch(db, kb.id)
    db.commit()
    result = {"knowledge_count": len(chunks), "document_id": doc.id, "splitting": splitting}
    return result
